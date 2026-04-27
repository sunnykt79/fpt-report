import io
from docx import Document
from docx.shared import Inches

def create_word_report(metrics_dict, start_date, end_date, figs=None):
    doc = Document()
    doc.add_heading('Báo Cáo Bán Hàng FPT Telecom', 0)
    
    doc.add_paragraph(f'Thời gian báo cáo: Từ {start_date} đến {end_date}')
    
    doc.add_heading('Tóm tắt Chỉ Số (KPIs)', level=1)
    
    p = doc.add_paragraph()
    p.add_run('Tổng PTTB: ').bold = True
    p.add_run(f"{metrics_dict['total_pttb']}\n")
    
    p.add_run('PTTB Combo: ').bold = True
    p.add_run(f"{metrics_dict['combo']}\n")
    
    p.add_run('PTTB V.Vip/Vip: ').bold = True
    p.add_run(f"{metrics_dict['vvip']}\n")
    
    p.add_run('Tỷ lệ Online: ').bold = True
    p.add_run(f"{metrics_dict['online_rate']:.2f}%\n")
    
    p.add_run('Tỷ lệ Hủy TSD: ').bold = True
    p.add_run(f"{metrics_dict['cancel_rate']:.2f}%\n")
    
    p.add_run('Thời gian triển khai trung bình: ').bold = True
    p.add_run(f"{metrics_dict['avg_deploy_days']:.1f} ngày\n")
    
    doc.add_heading('Nhận xét', level=1)
    doc.add_paragraph('Báo cáo được tạo tự động từ Hệ thống FSRD.')
    
    if figs:
        doc.add_heading('Biểu đồ Phân tích', level=1)
        for fig in figs:
            img_bytes = fig.to_image(format="png")
            img_stream = io.BytesIO(img_bytes)
            doc.add_picture(img_stream, width=Inches(6.0))
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
