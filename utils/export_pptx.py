import io
from pptx import Presentation
from pptx.util import Inches

def create_pptx_report(metrics_dict, start_date, end_date, figs=None):
    prs = Presentation()
    
    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "Báo Cáo Bán Hàng FPT Telecom"
    subtitle.text = f"Thời gian: Từ {start_date} đến {end_date}"
    
    # KPIs slide
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "Tóm Tắt Chỉ Số (KPIs)"
    
    tf = body_shape.text_frame
    tf.text = f"Tổng PTTB: {metrics_dict['total_pttb']}"
    
    p = tf.add_paragraph()
    p.text = f"PTTB Combo: {metrics_dict['combo']}"
    
    p = tf.add_paragraph()
    p.text = f"PTTB V.Vip/Vip: {metrics_dict['vvip']}"
    
    p = tf.add_paragraph()
    p.text = f"Tỷ lệ Online: {metrics_dict['online_rate']:.2f}%"
    
    p = tf.add_paragraph()
    p.text = f"Tỷ lệ Hủy TSD: {metrics_dict['cancel_rate']:.2f}%"
    
    p = tf.add_paragraph()
    p.text = f"Thời gian triển khai TB: {metrics_dict['avg_deploy_days']:.1f} ngày"
    
    if figs:
        for fig in figs:
            img_bytes = fig.to_image(format="png")
            img_stream = io.BytesIO(img_bytes)
            slide = prs.slides.add_slide(prs.slide_layouts[5]) # 5 is Title Only layout
            if slide.shapes.title:
                slide.shapes.title.text = "Biểu đồ Phân tích"
            slide.shapes.add_picture(img_stream, Inches(1), Inches(1.5), width=Inches(8))
    
    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer
